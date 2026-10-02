import os
try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
except ImportError:
    print("python-docx is not installed. Please install it.")
    exit(1)

def create_part1():
    doc = Document()
    title = doc.add_heading('EventLink CM Textbook - Part 1: UI/UX & Frontend Structure', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    doc.add_heading('Chapter 1: The Design Vision', level=1)
    doc.add_paragraph("This section of the textbook covers the visual layout and design structure of EventLink CM.")
    doc.add_paragraph("We overhauled the original design to feature a sleek, 'Apple-esque' aesthetic. We removed the old cyberpunk pink and purple colors, replacing them with a strict palette of white, dark gray, and vibrant blue. We implemented 'glassmorphism' (frosted glass effects) on the navigation bars and gave all buttons a rounded, pill-shaped design.")

    doc.add_heading('Chapter 2: HTML - The Skeleton', level=1)
    doc.add_paragraph("We structured 9 different pages (Home, About, Contact, Register, Login, Dashboard, Create Event, Organized Events, Analytics) using semantic HTML. HTML acts as the skeleton or walls of the website, using tags like <header>, <section>, and <nav> to group content logically.")

    doc.add_heading('Chapter 3: CSS - The Paint & Decoration', level=1)
    doc.add_paragraph("We wrote custom stylesheets (like futuristic.css) to apply our clean MacBook aesthetic across the entire platform. By using CSS Selectors, we targeted specific HTML classes to apply properties like background-color, border-radius (for the pill buttons), and box-shadows (to make the event cards 'float' off the page).")

    file_path = os.path.join('G:\\EventLink-CM_project', 'Textbook_Part1_UI_and_Design.docx')
    doc.save(file_path)

def create_part2():
    doc = Document()
    title = doc.add_heading('EventLink CM Textbook - Part 2: Frontend Logic & Interactivity', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    doc.add_heading('Chapter 1: JavaScript - The Electricity', level=1)
    doc.add_paragraph("This section of the textbook covers how we made the website interactive using JavaScript. If HTML is the walls and CSS is the paint, JavaScript is the electricity.")
    
    doc.add_heading('Chapter 2: Smart Routing & Memory', level=1)
    doc.add_paragraph("We built a 'smart routing' system that remembers what a user was trying to do. If a logged-out user clicks 'Host an Event', they aren't shown an ugly error popup. Instead, our JavaScript uses localStorage (the browser's memory) to save their intention. The code then instantly teleports the user to the Registration page. Once they finish signing up, the system checks the memory, silently finishes their original action, and sends them to the correct page.")

    doc.add_heading('Chapter 3: The Fetch API (Talking to the Server)', level=1)
    doc.add_paragraph("JavaScript uses asynchronous 'Fetch' requests to securely send user data (like login passwords or event RSVP requests) over the internet to the backend. It uses 'JSON' (JavaScript Object Notation) as the universal language to package this data so the Python server can understand it.")

    file_path = os.path.join('G:\\EventLink-CM_project', 'Textbook_Part2_Frontend_Logic.docx')
    doc.save(file_path)

def create_part3():
    doc = Document()
    title = doc.add_heading('EventLink CM Textbook - Part 3: Backend, API & Database', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    doc.add_heading('Chapter 1: The Python Backend', level=1)
    doc.add_paragraph("This section of the textbook covers the 'Brain' of the application. We built a Python backend using a framework called FastAPI. It listens for internet traffic and serves the web pages.")

    doc.add_heading('Chapter 2: Database Models & OOP', level=1)
    doc.add_paragraph("We use a SQLite database as our permanent filing cabinet. We designed our database using Object-Oriented Programming (OOP) in a file called models.py. A 'Class' acts as a blueprint, and an 'Object' is the actual data (like a specific User or Event ticket).")
    doc.add_paragraph("We also use 'Schemas' (schemas.py) which act as bouncers to validate incoming data from the internet before it is allowed into the database.")

    doc.add_heading('Chapter 3: The Core Algorithms', level=1)
    doc.add_paragraph("We implemented advanced algorithms for our Analytics page. When an event creator checks their dashboard, the Python backend executes a 'For-Loop'. It looks at every single RSVP ticket for that event, checks the user's gender, and adds +1 to the 'male' or 'female' math counters, then sends the final totals to the screen.")
    doc.add_paragraph("For the QR Scanner check-in, we use a Boolean (True/False) flag. When an attendee is scanned at the door, the backend flips scanned = True and permanently stamps the exact arrival time onto the ticket using the server's internal clock (datetime.now()).")

    file_path = os.path.join('G:\\EventLink-CM_project', 'Textbook_Part3_Backend_and_Database.docx')
    doc.save(file_path)

if __name__ == "__main__":
    create_part1()
    create_part2()
    create_part3()
    print("All three textbooks successfully created!")
