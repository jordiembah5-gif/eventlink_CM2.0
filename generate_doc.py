import os
try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
except ImportError:
    print("python-docx is not installed. Please install it.")
    exit(1)

def create_document():
    doc = Document()

    # Title
    title = doc.add_heading('EventLink CM: Project Journey & Logic Guide', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    doc.add_paragraph("This document outlines the step-by-step journey of building EventLink CM, a full-stack web application for discovering and hosting events in Cameroon. It is written to be easily understood by non-coders and beginners learning web development.")

    doc.add_heading('1. The Vision & Design', level=1)
    doc.add_paragraph("We started with a vision to build a clean, modern event platform. We overhauled the original design to feature a sleek, 'Apple-esque' aesthetic. We removed the old cyberpunk pink and purple colors, replacing them with a strict palette of white, dark gray, and vibrant blue. We implemented 'glassmorphism' (frosted glass effects) on the navigation bars and gave all buttons a rounded, pill-shaped design.")

    doc.add_heading('2. The Frontend (The User Interface)', level=1)
    doc.add_paragraph("We separated our frontend into three core technologies:")
    doc.add_paragraph("• HTML (The Skeleton): We structured 9 different pages (Home, About, Contact, Register, Login, Dashboard, Create Event, Organized Events, Analytics) using semantic HTML.", style='List Bullet')
    doc.add_paragraph("• CSS (The Paint): We wrote custom stylesheets (like futuristic.css) to apply our clean MacBook aesthetic across the entire platform.", style='List Bullet')
    doc.add_paragraph("• JavaScript (The Electricity): We added interactivity. We built a 'smart routing' system that remembers what a user was trying to do, sends them to the registration page, and automatically finishes their action once they are signed in, without showing any annoying popups.", style='List Bullet')

    doc.add_heading('3. The Backend & Database (The Brain)', level=1)
    doc.add_paragraph("To make the application real, we built a Python backend using a framework called FastAPI. We recently refactored this backend into a textbook-style format to make it easy to learn from:")
    doc.add_paragraph("• database.py: Connects our app to a SQLite database (our permanent filing cabinet).", style='List Bullet')
    doc.add_paragraph("• models.py: Defines our database tables using Object-Oriented Programming (OOP) blueprints for Users, Events, and RSVPs.", style='List Bullet')
    doc.add_paragraph("• schemas.py: Acts as a 'bouncer' to validate incoming data from the internet before it is allowed into the database.", style='List Bullet')
    doc.add_paragraph("• main.py: The central brain that handles API routes (communication channels) and serves our HTML pages to the user.", style='List Bullet')

    doc.add_heading('4. The Core Logic & Algorithms Explained', level=1)
    doc.add_paragraph("Here is how the hidden business logic works behind the scenes:")
    
    doc.add_heading('A. The Smart Routing Logic', level=2)
    doc.add_paragraph("When a logged-out user clicks 'Join Event', they aren't shown an ugly error popup. Instead, our JavaScript uses localStorage (the browser's memory) to save the phrase: 'pendingRSVP = event-name'. The code then instantly teleports the user to the Registration page. Once they finish signing up, the system checks the memory, sees the pending RSVP, silently tells the Python backend to register them for the event, and then sends them to their personalized dashboard. It feels like magic to the user!")
    
    doc.add_heading('B. The Analytics & Loop Logic', level=2)
    doc.add_paragraph("When an event creator checks their Analytics page, they see exactly how many males and females are attending. How does the code do this? The Python backend executes a 'For-Loop'. It looks at every single RSVP ticket for that event. For each ticket, it finds the user's account, checks their gender, and adds +1 to the 'male' or 'female' math counters, then sends the final totals to the screen in real time.")
    
    doc.add_heading('C. The QR Scanner Check-in Logic', level=2)
    doc.add_paragraph("Every digital ticket has a hidden property called a 'Boolean' (which means True or False). By default, scanned = False. When the organizer clicks 'Mock Scan QR' at the door, the Javascript sends a message to the backend. The backend flips scanned = True, looks at the server's internal clock (datetime.now()), and permanently stamps the exact arrival time onto the ticket in the database.")

    doc.add_heading('5. Deployment Preparation', level=1)
    doc.add_paragraph("Finally, we ran comprehensive integration tests across the entire platform. We verified that the database correctly stores information and the API endpoints respond accurately. We then generated a 'requirements.txt' file, ensuring the platform is 100% ready to be deployed live to the internet using cloud services like Render.")

    file_path = os.path.join('G:\\EventLink-CM_project', 'EventLink_CM_Project_Journey.docx')
    doc.save(file_path)
    print(f"Document successfully created at {file_path}")

if __name__ == "__main__":
    create_document()
