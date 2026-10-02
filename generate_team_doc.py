import os
try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
except ImportError:
    print("python-docx is not installed.")
    exit(1)

md_path = r"C:\Users\jordi\.gemini\antigravity\brain\2d3814a1-29f2-49a2-8be2-829dbd744638\EventLink_Team_Textbook.md"
docx_path = r"G:\EventLink-CM_project\EventLink_Team_Textbook_Render_Guide.docx"

def convert_md_to_docx():
    doc = Document()
    
    try:
        with open(md_path, 'r', encoding='utf-8') as f:
            lines = f.readlines()
    except FileNotFoundError:
        print("Markdown file not found.")
        return

    in_code_block = False
    
    for line in lines:
        raw_line = line.strip('\n')
        stripped = raw_line.strip()
        
        if stripped.startswith("```"):
            in_code_block = not in_code_block
            if in_code_block:
                doc.add_paragraph()
            continue
            
        if in_code_block:
            p = doc.add_paragraph(raw_line)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.left_indent = Pt(20)
            for run in p.runs:
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0, 102, 0)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith("# "):
            title = doc.add_heading(stripped[2:], level=0)
            title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        elif stripped == "":
            continue 
        else:
            doc.add_paragraph(stripped)

    doc.save(docx_path)
    print(f"Successfully generated Word Document at: {docx_path}")

if __name__ == "__main__":
    convert_md_to_docx()
